# Company: EonHive Inc. | Title: DOCX Conversion Fixture Builder
# Purpose: Create a real Word source with semantic styles, lists, links, a table and an image.
# Author: Stan Nesi | Created: 2026-10-27 | Updated: 2026-10-27
# Notes: Vibe coded with Codex.
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

root = Path(__file__).resolve().parent
document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.8)
for name in ["Normal", "Title", "Heading 1", "Heading 2", "Header", "Footer"]:
    style = document.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_after = Pt(6)
for style in document.styles:
    properties = style.element.find(qn("w:pPr"))
    if properties is not None:
        for border in list(properties.findall(qn("w:pBdr"))):
            properties.remove(border)
document.styles["Normal"].font.size = Pt(11)
document.styles["Title"].font.size = Pt(24)
document.styles["Heading 1"].font.size = Pt(15)
document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
section.header.paragraphs[0].text = "PRD Convert example"
section.footer.paragraphs[0].text = "Sample source document"
document.core_properties.title = "Portable Word report"
document.core_properties.author = "Stan Nesi"
document.core_properties.created = document.core_properties.modified = datetime(2026, 10, 27, tzinfo=timezone.utc)
document.add_paragraph("Portable Word report", "Title")
document.add_paragraph("This source demonstrates Word content that PRD Convert can reconstruct for responsive reading. Original layout and unsupported parts remain available in the retained source.")
document.add_heading("Structured content", 1)
paragraph = document.add_paragraph("Paragraphs retain their readable text. ")
paragraph.add_run("Emphasis").bold = True
paragraph.add_run(" becomes plain text in the canonical content model.")
for text in ["Keep useful document structure", "Report conversion limitations"]:
    document.add_paragraph(text, "List Bullet")
for text in ["Extract the document", "Validate and package the result"]:
    document.add_paragraph(text, "List Number")
heading = document.add_heading("Conversion details", 1)
bookmark = OxmlElement("w:bookmarkStart")
bookmark.set(qn("w:id"), "1")
bookmark.set(qn("w:name"), "details")
heading._p.insert(0, bookmark)
end = OxmlElement("w:bookmarkEnd")
end.set(qn("w:id"), "1")
heading._p.append(end)
table = document.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width, table.columns[1].width = Inches(2.1), Inches(4.8)
table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Source feature", "PRD result"
header = OxmlElement("w:tblHeader")
table.rows[0]._tr.get_or_add_trPr().append(header)
for values in [("Heading styles", "Structured heading nodes"), ("Declared image", "Portable raster asset"), ("Page layout", "Reported as unsupported")]:
    cells = table.add_row().cells
    for cell, text in zip(cells, values):
        cell.text = text
for row_index, row in enumerate(table.rows):
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "203B5A" if row_index == 0 else "F2F5F8" if row_index % 2 else "FFFFFF")
        properties.append(shading)
        margins = OxmlElement("w:tcMar")
        for side in ["top", "left", "bottom", "right"]:
            margin = OxmlElement(f"w:{side}")
            margin.set(qn("w:w"), "90")
            margin.set(qn("w:type"), "dxa")
            margins.append(margin)
        properties.append(margins)
        borders = OxmlElement("w:tcBorders")
        for side in ["top", "left", "bottom", "right"]:
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        properties.append(borders)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)
                run.bold = row_index == 0
picture_paragraph = document.add_paragraph()
picture_paragraph.paragraph_format.space_before = Pt(8)
picture = picture_paragraph.add_run().add_picture(str(root.parent / "conversion-markdown/flow.png"), width=Inches(4.7))
picture._inline.docPr.set("descr", "Source content passes through conversion into a portable PRD package.")
paragraph = document.add_paragraph()
hyperlink = OxmlElement("w:hyperlink")
hyperlink.set(qn("w:anchor"), "details")
run, text = OxmlElement("w:r"), OxmlElement("w:t")
text.text = "Return to conversion details"
run.append(text)
hyperlink.append(run)
paragraph._p.append(hyperlink)
paragraph = document.add_paragraph()
hyperlink = OxmlElement("w:hyperlink")
hyperlink.set(qn("r:id"), paragraph.part.relate_to("https://learn.microsoft.com/en-us/office/open-xml/word/structure-of-a-wordprocessingml-document", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
run, text = OxmlElement("w:r"), OxmlElement("w:t")
text.text = "WordprocessingML structure reference"
run.append(text)
hyperlink.append(run)
paragraph._p.append(hyperlink)

snapshot = BytesIO()
document.save(snapshot)
with ZipFile(snapshot) as source, ZipFile(root / "source.docx", "w") as output:
    for name in sorted(source.namelist()):
        info = ZipInfo(name, (1980, 1, 1, 0, 0, 0))
        info.compress_type = ZIP_DEFLATED
        output.writestr(info, source.read(name))
